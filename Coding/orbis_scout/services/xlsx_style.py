"""xlsx_style.py — shared dark-theme cell styling for openpyxl exports.

Extracted so multiple export endpoints (scatter, playing-styles, …) render
with one consistent look: dark header, zebra rows, tier colours, thin borders.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Palette ───────────────────────────────────────────────────────────────────
C_BG0   = '0d1117'
C_BG1   = '0f172a'
C_HDR   = '1e293b'
C_TITLE = '0d1b2a'
C_TEXT  = 'E2E8F0'
C_TEXT2 = '94A3B8'
C_ACCENT = '38BDF8'
C_TIER  = {
    'Elite': 'F59E0B', 'Advanced': '34D399', 'Good': '60A5FA',
    'Average': 'CBD5E1', 'Below Average': 'F87171',
}

_SIDE = Side(style='thin', color='1e293b')
_THIN = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)


def bg(r):
    """Zebra background for a 1-indexed row."""
    return C_BG0 if r % 2 == 0 else C_BG1


def hdr_cell(ws, row, col, value, fg=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(bold=True, color=C_TEXT, size=9, name='Consolas')
    c.fill      = PatternFill('solid', fgColor=fg or C_HDR)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border    = _THIN
    return c


def data_cell(ws, row, col, value, align='center', color=C_TEXT, bold=False, row_num=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(size=9, color=color, name='Consolas', bold=bold)
    c.fill      = PatternFill('solid', fgColor=bg(row_num or row))
    c.alignment = Alignment(horizontal=align, vertical='center')
    c.border    = _THIN
    return c


def tier_cell(ws, row, col, tier_str, row_num=None):
    c = ws.cell(row=row, column=col, value=tier_str)
    c.font      = Font(size=9, color=C_TIER.get(tier_str, C_TEXT2), name='Consolas', bold=True)
    c.fill      = PatternFill('solid', fgColor=bg(row_num or row))
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border    = _THIN
    return c


def title_row(ws, text, col_span, row=1):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    c = ws.cell(row=row, column=1, value=text)
    c.font      = Font(bold=True, size=13, color=C_TEXT, name='Consolas')
    c.fill      = PatternFill('solid', fgColor=C_TITLE)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 24


def meta_row(ws, text, col_span, row=2):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    c = ws.cell(row=row, column=1, value=text)
    c.font      = Font(size=8, color=C_TEXT2, name='Consolas')
    c.fill      = PatternFill('solid', fgColor=C_HDR)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 14


def tier_for(pct):
    """Percentile (0-100) → tier label."""
    if pct is None:    return '—'
    if pct >= 90:      return 'Elite'
    if pct >= 75:      return 'Advanced'
    if pct >= 50:      return 'Good'
    if pct >= 25:      return 'Average'
    return 'Below Average'
